"""Build the Word manuscript from the Markdown draft exported from the shared document.

- Every "[TO FILL: …]" and "[VERIFY: …]" span is highlighted in yellow.
- LaTeX display equations are typeset as images (matplotlib mathtext) and numbered.
- Figures from ../figures are placed above their captions.
- A4, Times New Roman 12 pt, 1.5 line spacing, continuous line numbers (for review).

    pip install python-docx matplotlib
    python paper/manuscript/build_docx.py      # writes paper/manuscript/manuscript_draft.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

HERE = Path(__file__).resolve().parent
FIG = HERE.parent / "figures"
SRC = HERE / "manuscript_draft.md"
OUT = HERE / "manuscript_draft.docx"

GAP = re.compile(r"\[(?:TO FILL|VERIFY)[^\]]*\]")
FIGFILES = {"Fig. 1.": "Fig1_worked_example", "Fig. 2.": "Fig2_global_recipes", "Fig. 3.": "Fig3_grades_vs_shape",
            "Fig. 4.": "Fig4_beta_scan", "Fig. 5.": "Fig5_robustness", "Fig. 6.": "Fig6_environment",
            "Fig. S1.": "FigS1_validation_demo"}
# mathtext versions of the display equations, in order of appearance
EQUATIONS = [
    r"$F(\tau) = f_0' + (1-f_0')\,[\,1-\exp(-(\tau/\lambda)^{\beta})\,]$",
    r"$\lambda = D\,/\,[\,\ln(\,(1-f_0')/0.2\,)\,]^{1/\beta}$",
    r"$\frac{d\tau}{dt} = f_{soil}\,a_P\,a_S\,[\,1-s\,x(\psi)\,]\;"
    r"\overline{\exp[\,\frac{E_a}{R}(\frac{1}{298.15}-\frac{1}{T+273.15})\,]\,r(T)}$",
    r"$P(t+1) = P(t)\,(1-k) + R(t) - U(t)/\eta$",
    r"$\min\ \sum_k c_k x_k\quad \mathrm{s.t.}\quad \sum_k x_k\,G_k^{(\Delta T)}(t) \geq N^{(\Delta T)}(t)"
    r"\quad \forall t,\ \Delta T \in \{-1.5,\,0,\,+1.5\}$",
]


def unescape(s):
    return re.sub(r"\\([\[\]\*_`#\\\-\.\(\)!>|])", r"\1", s)


def parse_inline(text):
    """Markdown inline → list of (text, bold, italic, code). Links keep their text only."""
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1", text)
    out = []
    pat = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))")
    pos = 0
    for m in pat.finditer(text):
        if m.start() > pos:
            out.append((unescape(text[pos:m.start()]), False, False, False))
        tok = m.group(0)
        if tok.startswith("**"):
            for t, _, it, cd in parse_inline(tok[2:-2]):
                out.append((t, True, it, cd))
        elif tok.startswith("`"):
            out.append((tok[1:-1], False, False, True))
        else:
            out.append((unescape(tok[1:-1]), False, True, False))
        pos = m.end()
    if pos < len(text):
        out.append((unescape(text[pos:]), False, False, False))
    return out


def add_runs(par, text, size=None):
    segs = parse_inline(text)
    full = "".join(s[0] for s in segs)
    marks = [False] * len(full)
    for m in GAP.finditer(full):
        for i in range(m.start(), m.end()):
            marks[i] = True
    i = 0
    for t, b, it, cd in segs:
        j = 0
        while j < len(t):
            k = j
            while k < len(t) and marks[i + k] == marks[i + j]:
                k += 1
            r = par.add_run(t[j:k])
            r.bold, r.italic = b, it
            if cd:
                r.font.name = "Courier New"
            if size:
                r.font.size = size
            if marks[i + j]:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            j = k
        i += len(t)


def equation_image(tex, idx):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig = plt.figure(figsize=(0.01, 0.01))
    fig.text(0, 0, tex, fontsize=13, math_fontfamily="stix")
    path = HERE / f"_eq{idx}.png"
    fig.savefig(path, dpi=300, bbox_inches="tight", pad_inches=0.04, transparent=False, facecolor="white")
    plt.close(fig)
    return path


def line_numbers(section):
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    ln.set(qn("w:distance"), "283")
    section._sectPr.append(ln)


def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        for j, c in enumerate(r):
            p = t.cell(i, j).paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            add_runs(p, f"**{c}**" if i == 0 and not c.startswith("**") else c, size=Pt(9))
    doc.add_paragraph()


def build():
    md = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(25))
    line_numbers(sec)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(12)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 14), ("Heading 2", 13), ("Heading 3", 12)):
        h = doc.styles[name]
        h.font.name, h.font.size, h.font.bold = "Times New Roman", Pt(size), True
        h.font.color.rgb = None
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    i, eq, in_captions = 0, 0, False
    # skip the document title and byline of the shared doc
    while i < len(md) and (md[i].startswith("# ") or md[i].startswith("Oct ") or not md[i].strip()):
        i += 1
    while i < len(md):
        line = md[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            body = []
            i += 1
            while i < len(md) and not md[i].startswith("```"):
                body.append(md[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if eq < len(EQUATIONS):
                img = equation_image(EQUATIONS[eq], eq)
                from PIL import Image
                w, h = Image.open(img).size
                height = Mm(9 if eq != 2 else 12)
                if height * w / h > Mm(140):           # keep wide equations inside the text block
                    p.add_run().add_picture(str(img), width=Mm(140))
                else:
                    p.add_run().add_picture(str(img), height=height)
                p.add_run(f"\t({eq + 1})")
            else:
                add_runs(p, "[VERIFY: equation could not be typeset] " + " ".join(body))
            eq += 1
            continue
        m = re.match(r"^(#{2,3}) (.*)", line)
        if m:
            level = len(m.group(1)) - 1
            doc.add_heading(unescape(m.group(2)), level=level)
            in_captions = m.group(2).startswith("Figure captions")
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(md) and md[i].startswith("|"):
                rows.append(md[i])
                i += 1
            table(doc, rows)
            continue
        m = re.match(r"^(\s*)(- |\d+\. )(.*)", line)
        if m:
            style = "List Bullet" if m.group(2) == "- " else "List Number"
            p = doc.add_paragraph(style=style)
            add_runs(p, m.group(3))
            i += 1
            continue
        if in_captions:
            key = next((k for k in FIGFILES if line.startswith(f"**{k}**")), None)
            if key:
                img = FIG / f"{FIGFILES[key]}.png"
                if img.exists():
                    doc.add_picture(str(img), width=Mm(160))
                    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        p = doc.add_paragraph()
        add_runs(p, line)
        i += 1
    doc.save(OUT)
    for f in HERE.glob("_eq*.png"):
        f.unlink()
    return OUT


if __name__ == "__main__":
    out = build()
    from docx import Document as D
    d = D(out)
    n_hl = sum(1 for p in d.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
    n_hl += sum(1 for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs for r in p.runs
                if r.font.highlight_color == WD_COLOR_INDEX.YELLOW)
    print(f"wrote {out} ({out.stat().st_size // 1024} kB), yellow runs: {n_hl}")
    sys.exit(0)
